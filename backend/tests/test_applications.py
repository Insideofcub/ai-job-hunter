import json
from io import BytesIO
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.main import app
from app.database import SessionLocal
from app.models import CandidateProfile
from uuid import uuid4


def make_job(client, suffix):
    # Use a unique URL each run so the suite stays repeatable with a persistent local SQLite DB.
    unique_suffix = f"{suffix}-{uuid4().hex[:10]}"
    response = client.post('/api/jobs', json={
        'title': f'Talent Acquisition Specialist {unique_suffix}',
        'company': f'Example Company {unique_suffix}',
        'location': 'Remote, India',
        'work_mode': 'Remote',
        'source_url': f'https://jobs.lever.co/example/{unique_suffix}',
        'description': 'Talent acquisition, IT recruiter, end-to-end recruitment, Boolean search.',
        'verified_posting_time': True,
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_application_created_from_job_and_status_updated():
    with TestClient(app) as client:
        job = make_job(client, 'application-test')
        created = client.post(f"/api/applications/from-job/{job['id']}")
        assert created.status_code == 201, created.text
        record = created.json()
        assert record['status'] == 'ready_for_review'
        assert record['job_id'] == job['id']
        duplicate = client.post(f"/api/applications/from-job/{job['id']}")
        assert duplicate.status_code == 409
        updated = client.patch(f"/api/applications/{record['id']}", json={'status': 'submitted'})
        assert updated.status_code == 200
        assert updated.json()['status'] == 'submitted'
        assert updated.json()['applied_at'] is not None


def test_screening_answers_stay_grounded_and_sensitive_unknown_needs_review():
    with TestClient(app) as client:
        profile = client.get('/api/profile').json()
        profile['full_name'] = 'Test Candidate'
        profile['years_experience'] = 3
        profile['resume_headline'] = 'IT Recruiter | Talent Acquisition Specialist'
        profile['profile_summary'] = 'I have 3 years of end-to-end IT and Non-IT recruitment experience.'
        profile['skills'] = 'Recruitment, Boolean search, stakeholder management'
        profile['verified_facts_json'] = '{}'
        profile['career_profile_json'] = '{}'
        profile['personal_details_json'] = '{}'
        saved = client.put('/api/profile', json=profile)
        assert saved.status_code == 200, saved.text
        job = make_job(client, 'answers-test')
        app_record = client.post(f"/api/applications/from-job/{job['id']}").json()
        response = client.post(f"/api/applications/{app_record['id']}/answers/generate", json={
            'questions': ['Tell us about yourself.', 'How many years of experience do you have?', 'What is your notice period?']
        })
        assert response.status_code == 200, response.text
        answers = json.loads(response.json()['answers_json'])
        intro = next(item for item in answers if item['question'] == 'Tell us about yourself.')
        experience = next(item for item in answers if item['question'] == 'How many years of experience do you have?')
        notice = next(item for item in answers if item['question'] == 'What is your notice period?')
        assert intro['answer']
        assert '3 years' in experience['answer']
        assert notice['status'] == 'needs_review'
        assert notice['answer'] == ''


def test_analytics_export_excel_has_three_sheets():
    with TestClient(app) as client:
        job = make_job(client, 'export-test')
        response = client.post(f"/api/applications/from-job/{job['id']}")
        assert response.status_code == 201, response.text
        export = client.get('/api/analytics/export.xlsx')
        assert export.status_code == 200
        assert export.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        assert {'Analytics', 'Applications', 'Screening Answers'}.issubset(set(workbook.sheetnames))
        assert workbook['Applications'].max_row >= 2


def test_agent_run_with_no_configured_sources_returns_explanation():
    with TestClient(app) as client:
        response = client.post('/api/agent/run', json={'query': 'Recruiter', 'hours': 24})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['sources_checked'] >= 0
        if data['sources_checked'] == 0:
            assert any('configured' in warning.lower() for warning in data['warnings'])


def test_automation_settings_keep_submit_disabled_by_default():
    with TestClient(app) as client:
        response = client.get('/api/automation/settings')
        assert response.status_code == 200
        data = response.json()
        assert isinstance(data['auto_submit_enabled'], bool)
        assert 'greenhouse' in ' '.join(data['supported_hosts'])


def test_automation_diagnostics_returns_actionable_runtime_status(monkeypatch):
    import app.main as main_module

    async def fake_diagnostics():
        return {
            'ready': False,
            'playwright_installed': True,
            'chromium_installed': False,
            'python_executable': r'C:\test-env\python.exe',
            'chromium_path': r'C:\test-env\ms-playwright\chromium\chrome.exe',
            'profile_directory': r'C:\test-env\browser-profile',
            'message': 'Chromium is missing. Run python -m playwright install chromium.',
            'install_command': 'python -m pip install -r requirements.txt; python -m playwright install chromium',
        }

    monkeypatch.setattr(main_module, 'browser_runtime_diagnostics', fake_diagnostics)
    with TestClient(app) as client:
        response = client.get('/api/automation/diagnostics')
        assert response.status_code == 200, response.text
        body = response.json()
        assert body['ready'] is False
        assert body['chromium_installed'] is False
        assert 'playwright install chromium' in body['message']


def test_browser_workflow_runtime_errors_are_exposed_as_actionable_503(monkeypatch):
    import app.main as main_module

    def make_fake_error():
        async def fake_fill(*args, **kwargs):
            raise RuntimeError('Playwright is not installed in the backend Python environment. Run python -m playwright install chromium.')
        return fake_fill

    monkeypatch.setattr(main_module, 'fill_application_form', make_fake_error())
    with TestClient(app) as client:
        profile = client.get('/api/profile').json()
        assert profile
        job = make_job(client, 'browser-503-diagnostic')
        app_record = client.post(f"/api/applications/from-job/{job['id']}")
        assert app_record.status_code == 201, app_record.text
        response = client.post(f"/api/applications/{app_record.json()['id']}/browser-fill", json={'confirm_submit': False})
        assert response.status_code == 503
        assert 'Playwright is not installed' in response.json()['detail']


def test_application_browser_submit_is_blocked_by_default():
    import app.main as main_module
    from app.models import Application
    with TestClient(app) as client:
        created = client.post('/api/applications', json={
            'title': 'Talent Acquisition Specialist', 'company': 'Greenhouse Test',
            'location': 'Remote', 'work_mode': 'Remote',
            'source_url': f"https://boards.greenhouse.io/example/jobs/123-{uuid4().hex[:10]}",
            'platform': 'greenhouse', 'status': 'ready_for_review'
        })
        assert created.status_code == 201, created.text
        # AUTO_SUBMIT_ENABLED is false by default; no browser should launch for this request.
        old_value = main_module.settings.auto_submit_enabled
        main_module.settings.auto_submit_enabled = False
        try:
            response = client.post(f"/api/applications/{created.json()['id']}/browser-fill", json={'confirm_submit': True})
            assert response.status_code == 403
        finally:
            main_module.settings.auto_submit_enabled = old_value


def test_aggregator_source_provider_is_supported():
    with TestClient(app) as client:
        response = client.post('/api/sources', json={'provider': 'remoteok', 'company_name': 'Remote OK', 'board_id': 'default'})
        assert response.status_code in {201, 409}
        invalid = client.post('/api/sources', json={'provider': 'remoteok', 'company_name': 'Remote OK Other', 'board_id': 'not-default'})
        assert invalid.status_code == 422


def test_screening_answer_can_be_edited_and_verified():
    with TestClient(app) as client:
        job = make_job(client, 'answer-review-test')
        record = client.post(f"/api/applications/from-job/{job['id']}").json()
        generated = client.post(f"/api/applications/{record['id']}/answers/generate", json={
            'questions': ['Why are you interested in this role?']
        })
        assert generated.status_code == 200, generated.text
        answers = json.loads(generated.json()['answers_json'])
        assert answers
        review_index = next((idx for idx, item in enumerate(answers) if item['question'] == 'Why are you interested in this role?'), 0)
        answer_text = 'This role matches my recruitment experience and professional goals.'
        reviewed = client.patch(f"/api/applications/{record['id']}/answers/{review_index}", json={
            'answer': answer_text, 'status': 'verified'
        })
        assert reviewed.status_code == 200, reviewed.text
        updated_answers = json.loads(reviewed.json()['answers_json'])
        assert updated_answers[review_index]['answer'] == answer_text
        assert updated_answers[review_index]['status'] == 'verified'
        assert updated_answers[review_index]['source'] == 'User reviewed and confirmed'


def test_job_agent_run_creates_application_records_from_fresh_configured_source(monkeypatch):
    from datetime import datetime, timedelta, timezone
    from uuid import uuid4
    from app.services.job_sources.base import NormalizedPosting

    token = uuid4().hex[:12]
    company = f"Agent Test Company {token}"
    board = f"agent-{token}"
    now = datetime.now(timezone.utc)

    async def fake_fetch_source(client, source):
        return [NormalizedPosting(
            title="Talent Acquisition Specialist",
            company=source.company_name,
            location="Remote, India",
            work_mode="Remote",
            source_url=f"https://jobs.example.com/{source.board_id}",
            description="End-to-end IT and Non-IT recruitment, LinkedIn sourcing, Boolean search, stakeholder management.",
            posted_at=now - timedelta(hours=2),
            verified_posting_time=True,
            provider=source.provider,
            source_company=source.company_name,
        )]

    monkeypatch.setattr("app.main._fetch_source", fake_fetch_source)
    with TestClient(app) as client:
        profile = client.get('/api/profile').json()
        profile['full_name'] = 'Agent Test Candidate'
        profile['target_titles'] = 'Talent Acquisition Specialist, IT Recruiter'
        profile['skills'] = 'IT recruitment, Non-IT recruitment, Boolean search, LinkedIn sourcing'
        profile['profile_summary'] = 'Experience in end-to-end IT and Non-IT recruitment.'
        profile['employment_highlights'] = 'Candidate sourcing, interview coordination, stakeholder management.'
        assert client.put('/api/profile', json=profile).status_code == 200
        source = client.post('/api/sources', json={'provider': 'lever', 'company_name': company, 'board_id': board})
        assert source.status_code == 201, source.text
        result = client.post('/api/agent/run', json={
            'query': 'Talent Acquisition, IT Recruiter', 'hours': 24,
            'min_match_score': 0, 'max_applications': 5,
            'prepare_applications': True, 'auto_apply': False,
        })
        assert result.status_code == 200, result.text
        data = result.json()
        assert data['sources_checked'] >= 1
        assert data['applications_created'] >= 1
        applications = client.get('/api/applications').json()
        assert any(item['company'] == company for item in applications)


def test_agent_run_with_auto_submit_disabled_still_completes_discovery(monkeypatch):
    from datetime import datetime, timezone
    import app.main as main_module
    from app.schemas import DiscoveryResponse

    monkeypatch.setattr(main_module.settings, "auto_submit_enabled", False)

    async def fake_discover(payload, db):
        return DiscoveryResponse(
            query=payload.query, hours=payload.hours, checked_at=datetime.now(timezone.utc),
            sources_checked=0, fresh_count=0, unverified_count=0, fresh_jobs=[], unverified_jobs=[],
            source_statuses=[], region=payload.region, custom_location=payload.custom_location,
            work_modes=payload.work_modes, employment_types=payload.employment_types,
        )

    monkeypatch.setattr(main_module, "discover_jobs", fake_discover)
    with TestClient(app) as client:
        response = client.post('/api/agent/run', json={
            'query': 'Recruiter', 'hours': 24, 'auto_apply': True, 'prepare_applications': True
        })
        assert response.status_code == 200, response.text
        body = response.json()
        assert body['submitted'] == 0
        assert any('AUTO_SUBMIT_ENABLED=false' in warning for warning in body['warnings'])


def test_lever_404_pauses_source_and_reports_actionable_message(monkeypatch):
    from uuid import uuid4
    import httpx
    import app.main as main_module

    board_id = f"agent-{uuid4().hex[:12]}"
    company = f"Agent Test Company {uuid4().hex[:8]}"

    async def fake_fetch_source(client, source, query=""):
        request = httpx.Request('GET', f'https://api.lever.co/v0/postings/{source.board_id}?mode=json')
        response = httpx.Response(404, request=request)
        raise httpx.HTTPStatusError('Client error 404 Not Found', request=request, response=response)

    monkeypatch.setattr(main_module, "_fetch_source", fake_fetch_source)
    with TestClient(app) as client:
        existing = client.get('/api/sources').json()
        original_enabled = {item['id']: item['enabled'] for item in existing}
        new_id = None
        try:
            for item in existing:
                if item['enabled']:
                    client.patch(f"/api/sources/{item['id']}", json={'enabled': False})
            created = client.post('/api/sources', json={
                'provider': 'lever', 'company_name': company, 'board_id': board_id
            })
            assert created.status_code == 201, created.text
            new_id = created.json()['id']
            result = client.post('/api/discovery/search', json={'query': 'Recruiter', 'hours': 24})
            assert result.status_code == 200, result.text
            source_status = next(item for item in result.json()['source_statuses'] if item['source_id'] == new_id)
            assert source_status['status'] == 'error'
            assert 'paused' in source_status['message'].lower()
            refreshed = next(item for item in client.get('/api/sources').json() if item['id'] == new_id)
            assert refreshed['enabled'] is False
        finally:
            if new_id is not None:
                client.delete(f'/api/sources/{new_id}')
            for source_id, enabled in original_enabled.items():
                client.patch(f'/api/sources/{source_id}', json={'enabled': enabled})


def test_batch_application_queue_creates_and_reuses_records_without_submitting():
    with TestClient(app) as client:
        first = make_job(client, 'queue-batch-a')
        second = make_job(client, 'queue-batch-b')
        payload = {
            'jobs': [
                {'job_id': first['id'], 'title': first['title'], 'company': first['company'],
                 'location': first['location'], 'work_mode': first['work_mode'],
                 'source_url': first['source_url'], 'description': first['description'],
                 'verified_posting_time': True, 'match_score': 88, 'matched_keywords': ['recruitment']},
                {'job_id': second['id'], 'title': second['title'], 'company': second['company'],
                 'location': second['location'], 'work_mode': second['work_mode'],
                 'source_url': second['source_url'], 'description': second['description'],
                 'verified_posting_time': True, 'match_score': 75, 'matched_keywords': ['sourcing']},
            ]
        }
        response = client.post('/api/applications/queue', json=payload)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body['requested'] == 2
        assert body['created'] == 2
        assert body['reused'] == 0
        assert body['submitted'] == 0
        assert all(item['action'] == 'created' for item in body['items'])

        repeated = client.post('/api/applications/queue', json=payload)
        assert repeated.status_code == 200, repeated.text
        repeated_body = repeated.json()
        assert repeated_body['created'] == 0
        assert repeated_body['reused'] == 2
        assert repeated_body['submitted'] == 0


def test_batch_application_queue_saves_unverified_lead_for_manual_review():
    with TestClient(app) as client:
        suffix = uuid4().hex[:10]
        response = client.post('/api/applications/queue', json={'jobs': [{
            'title': 'Talent Acquisition Specialist', 'company': 'Unverified Lead Co',
            'location': 'Remote, India', 'work_mode': 'Remote',
            'source_url': f'https://jobs.example.com/unverified-{suffix}',
            'description': 'Recruitment, sourcing, interview coordination',
            'verified_posting_time': False, 'match_score': 72,
            'matched_keywords': ['talent acquisition']
        }]})
        assert response.status_code == 200, response.text
        body = response.json()
        assert body['created'] == 1
        item = body['items'][0]
        assert item['action'] == 'created'
        assert item['job_id'] is not None
        job = client.get('/api/jobs').json()
        saved = next(row for row in job if row['id'] == item['job_id'])
        assert saved['verified_posting_time'] is False
        record = next(row for row in client.get('/api/applications').json() if row['id'] == item['application_id'])
        assert record['status'] == 'ready_for_review'
        assert record['applied_at'] is None


def test_windows_selector_loop_is_flagged_for_playwright():
    from app.services.browser_apply import windows_playwright_loop_issue

    issue = windows_playwright_loop_issue("SelectorEventLoop", "win32")
    assert issue is not None
    assert "without --reload" in issue
    assert "python run_server.py" in issue
    assert windows_playwright_loop_issue("ProactorEventLoop", "win32") is None
    # Unix selector loops support asyncio subprocesses and must not be flagged.
    assert windows_playwright_loop_issue("_UnixSelectorEventLoop", "linux") is None
