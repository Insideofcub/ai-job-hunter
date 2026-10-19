import asyncio
import io
import json
from collections import Counter
import logging
import re
import zipfile
from datetime import timedelta
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from typing import Any
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
import httpx
from sqlalchemy import inspect, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .config import settings
from .database import Base, engine, get_db, SessionLocal
from .models import AppSetting, Application, CandidateProfile, Job, JobSource
from .schemas import (
    DiscoveryItem, DiscoveryRequest, DiscoveryResponse, HealthResponse, JobCreate,
    JobAnalysisResponse, JobMatchRead, JobRead, JobSourceCreate, JobSourceRead, JobSourceToggle, ProfileRead, ProfileUpdate,
    PhotoUploadResponse, ResumeTextResponse, ResumeUploadResponse, SourceCheckStatus,
    ApplicationCreate, ApplicationRead, ApplicationUpdate, ApplicationQueueRequest,
    ApplicationQueueResponse, ApplicationQueueItem, ScreeningAnswerRequest,
    AgentRunRequest, AgentRunResponse, BrowserApplyRequest, BrowserApplyResponse, ScreeningAnswerUpdate,
    AutomationSettingsResponse, AnalyticsResponse,
)
from .services.discovery import canonicalize_job_url, deduplicate_postings, filter_and_deduplicate, matches_query, query_match_reason, query_phrases
from .services.job_filters import classify_employment_type, classify_work_mode, location_matches_region
from .services.job_sources import fetch_greenhouse, fetch_lever
from .services.job_sources.jobspy import fetch_jobspy_postings, jobspy_package_available
from .services.job_sources.aggregators import COUNTRY_CODES, fetch_arbeitnow, fetch_remoteok, fetch_jobicy, fetch_himalayas, fetch_adzuna
from .services.job_sources.base import NormalizedPosting
from .services.application_workflow import (
    SUPPORTED_APPLICATION_HOSTS, draft_screening_answers, host_is_supported,
    parse_json as parse_workflow_json, platform_from_url, safe_resume_path,
)
from .services.browser_apply import fill_application_form, browser_runtime_diagnostics
from .services.ai_provider import draft_screening_answers_with_ai
from .services.web_research import fetch_brave_web_research
from .services.research_scoring import score_posting
from .services.job_analysis import analyze_job_fit
from .services.semantic_matching import semantic_scores_for_jobs

BACKEND_DIR = Path(__file__).resolve().parents[1]
RESUME_DIR = BACKEND_DIR / "data" / "resumes"
MAX_RESUME_BYTES = 2 * 1024 * 1024
logger = logging.getLogger("ai_job_hunter.discovery")
SUPPORTED_RESUME_SUFFIXES = {".pdf", ".doc", ".docx", ".rtf"}

PROFILE_COLUMN_DEFINITIONS = {
    "resume_headline": "VARCHAR(300) NOT NULL DEFAULT ''",
    "profile_summary": "TEXT NOT NULL DEFAULT ''",
    "employment_highlights": "TEXT NOT NULL DEFAULT ''",
    "employment_history_json": "TEXT NOT NULL DEFAULT '[]'",
    "education_json": "TEXT NOT NULL DEFAULT '[]'",
    "it_skills_json": "TEXT NOT NULL DEFAULT '[]'",
    "projects_json": "TEXT NOT NULL DEFAULT '[]'",
    "online_profiles_json": "TEXT NOT NULL DEFAULT '[]'",
    "work_samples_json": "TEXT NOT NULL DEFAULT '[]'",
    "publications_json": "TEXT NOT NULL DEFAULT '[]'",
    "presentations_json": "TEXT NOT NULL DEFAULT '[]'",
    "patents_json": "TEXT NOT NULL DEFAULT '[]'",
    "certifications_json": "TEXT NOT NULL DEFAULT '[]'",
    "career_profile_json": "TEXT NOT NULL DEFAULT '{}'",
    "personal_details_json": "TEXT NOT NULL DEFAULT '{}'",
    "languages_json": "TEXT NOT NULL DEFAULT '[]'",
    "resume_storage_key": "VARCHAR(100)",
    "resume_original_filename": "VARCHAR(255)",
    "resume_uploaded_at": "DATETIME",
    "resume_size_bytes": "INTEGER",
    "resume_parse_status": "VARCHAR(80) NOT NULL DEFAULT 'not_uploaded'",
    "resume_text": "TEXT NOT NULL DEFAULT ''",
    "photo_storage_key": "VARCHAR(100)",
    "photo_original_filename": "VARCHAR(255)",
    "photo_content_type": "VARCHAR(80)",
    "photo_size_bytes": "INTEGER",
    "photo_uploaded_at": "DATETIME",
}


def ensure_profile_schema_compatibility() -> None:
    """Add Phase 3 fields to older local SQLite databases without deleting data."""
    if not settings.database_url.startswith("sqlite"):
        return
    inspector = inspect(engine)
    if "candidate_profile" not in inspector.get_table_names():
        return
    existing = {column["name"] for column in inspector.get_columns("candidate_profile")}
    with engine.begin() as connection:
        for name, definition in PROFILE_COLUMN_DEFINITIONS.items():
            if name not in existing:
                connection.execute(text(f"ALTER TABLE candidate_profile ADD COLUMN {name} {definition}"))


def _extract_resume_text(data: bytes, suffix: str) -> tuple[str, str]:
    """Extract text when a supported parser is available; never run uploaded files."""
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=False)
        if len(reader.pages) > 40:
            raise HTTPException(status_code=413, detail="Resume PDF exceeds the 40-page safety limit.")
        text_parts = [(page.extract_text() or "") for page in reader.pages]
        extracted = "\n".join(text_parts).strip()
        return extracted[:200_000], "text_extracted" if extracted else "uploaded_no_selectable_text"
    if suffix == ".docx":
        if len(data) > MAX_RESUME_BYTES or not zipfile.is_zipfile(io.BytesIO(data)):
            raise HTTPException(status_code=400, detail="This DOCX file does not appear to be a valid Word document.")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if sum(item.file_size for item in infos) > 20 * 1024 * 1024 or len(infos) > 500:
                raise HTTPException(status_code=413, detail="The DOCX contents exceed the safe extraction limit.")
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(set(archive.namelist())):
                raise HTTPException(status_code=400, detail="This DOCX file does not appear to be a valid Word document.")
        from docx import Document
        document = Document(io.BytesIO(data))
        extracted = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        return extracted[:200_000], "text_extracted" if extracted else "uploaded_no_selectable_text"
    if suffix == ".rtf":
        if not data.lstrip().startswith(b"{\\rtf"):
            raise HTTPException(status_code=400, detail="This RTF file does not appear to be valid.")
        raw = data.decode("utf-8", errors="replace")
        # Basic, conservative RTF text extraction without executing embedded content.
        raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
        raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
        extracted = re.sub(r"[{}]", " ", raw)
        extracted = re.sub(r"\s+", " ", extracted).strip()
        return extracted[:200_000], "text_extracted" if extracted else "uploaded_no_selectable_text"
    # Legacy .doc files can be stored, but this starter does not include a binary DOC parser.
    if suffix == ".doc":
        if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise HTTPException(status_code=400, detail="This DOC file does not appear to be a valid legacy Word document.")
        return "", "uploaded_extraction_not_supported_for_legacy_doc"
    raise HTTPException(status_code=415, detail="Supported resume formats: PDF, DOC, DOCX, and RTF.")


def _safe_resume_path(storage_key: str | None) -> Path | None:
    if not storage_key or not re.fullmatch(r"[a-f0-9]{32}\.(pdf|doc|docx|rtf)", storage_key):
        return None
    path = (RESUME_DIR / storage_key).resolve()
    if path.parent != RESUME_DIR.resolve():
        return None
    return path


def _safe_photo_path(storage_key: str | None) -> Path | None:
    if not storage_key or not re.fullmatch(r"[a-f0-9]{32}\.(png|jpg|jpeg|webp)", storage_key):
        return None
    photo_dir = RESUME_DIR.parent / "photos"
    path = (photo_dir / storage_key).resolve()
    if path.parent != photo_dir.resolve():
        return None
    return path


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    ensure_profile_schema_compatibility()
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    # Keep a single local candidate profile in the local-first starter.
    with SessionLocal() as db:
        if db.get(CandidateProfile, 1) is None:
            db.add(CandidateProfile(id=1))
            db.commit()

        # Add useful no-key public aggregators once to existing local projects.
        # This corrects the common "only two company boards checked" setup while
        # allowing users to pause/remove a feed later without it reappearing.
        seed_key = "recommended_public_feeds_seeded_v3"
        if db.get(AppSetting, seed_key) is None:
            defaults = [
                ("arbeitnow", "default", "Arbeitnow Public Feed"),
                ("remoteok", "default", "Remote OK Public Feed"),
                ("jobicy", "default", "Jobicy Global + APAC Remote Feed"),
                ("himalayas", "default", "Himalayas Remote Jobs (attribution required)"),
            ]
            if settings.adzuna_app_id and settings.adzuna_app_key:
                defaults.extend([
                    ("adzuna", "in", "Adzuna India"),
                    ("adzuna", "gb", "Adzuna United Kingdom"),
                    ("adzuna", "us", "Adzuna United States"),
                ])
            for provider, board_id, company_name in defaults:
                existing = db.scalar(select(JobSource).where(
                    JobSource.provider == provider, JobSource.board_id == board_id
                ))
                if existing is None:
                    db.add(JobSource(provider=provider, board_id=board_id, company_name=company_name, enabled=True))
            db.add(AppSetting(key=seed_key, value="true"))
            db.commit()
        # Add the optional broad-web research provider once to existing Phase 4 databases.
        # It remains skipped until the user supplies BRAVE_SEARCH_API_KEY in backend/.env.
        brave_seed_key = "brave_search_source_seeded_v1"
        if db.get(AppSetting, brave_seed_key) is None:
            existing = db.scalar(select(JobSource).where(
                JobSource.provider == "brave", JobSource.board_id == "default"
            ))
            if existing is None:
                db.add(JobSource(
                    provider="brave", board_id="default", company_name="Brave Web Research",
                    enabled=True,
                ))
            db.add(AppSetting(key=brave_seed_key, value="true"))
            db.commit()
        # JobSpy is visible in source management but intentionally paused by default.
        # It requires a separate optional dependency and an explicit environment opt-in.
        jobspy_seed_key = "optional_jobspy_source_seeded_v1"
        if db.get(AppSetting, jobspy_seed_key) is None:
            existing = db.scalar(select(JobSource).where(
                JobSource.provider == "jobspy", JobSource.board_id == "default"
            ))
            if existing is None:
                db.add(JobSource(
                    provider="jobspy", board_id="default", company_name="JobSpy Multi-Portal Search",
                    enabled=False,
                ))
            db.add(AppSetting(key=jobspy_seed_key, value="true"))
            db.commit()
    yield


APP_VERSION = "0.6.3"
app = FastAPI(title=settings.app_name, version=APP_VERSION, lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origin_list,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
)


@app.get("/api/health", response_model=HealthResponse, tags=["system"])
def health():
    return {"status": "ok", "app": settings.app_name, "environment": settings.app_env, "version": APP_VERSION}


@app.get("/api/version", tags=["system"])
def version_info():
    """Return release metadata without exposing credentials or local paths."""
    return {"app": settings.app_name, "version": APP_VERSION, "environment": settings.app_env}


@app.get("/api/system/readiness", tags=["system"])
def system_readiness():
    """Local diagnostic summary. Reports configuration states, never secret values."""
    checks: list[dict[str, Any]] = []
    overall_ok = True
    try:
        with engine.connect() as connection:
            connection.execute(text("SELECT 1"))
        checks.append({"name": "database", "status": "ok", "message": "Database connection succeeded."})
    except Exception as exc:
        overall_ok = False
        checks.append({"name": "database", "status": "error", "message": f"Database connection failed: {type(exc).__name__}. Check DATABASE_URL and file permissions."})

    checks.append({
        "name": "resume_storage",
        "status": "ok" if RESUME_DIR.exists() and RESUME_DIR.is_dir() else "warning",
        "message": "Resume storage directory is available." if RESUME_DIR.exists() and RESUME_DIR.is_dir() else "Resume storage will be created during startup or upload.",
    })
    checks.append({
        "name": "ollama",
        "status": "configured" if settings.ollama_base_url and settings.ai_provider.lower() == "ollama" else "optional",
        "message": "Ollama is selected as the AI provider; verify the local service and model if AI-assisted features fail." if settings.ai_provider.lower() == "ollama" else "Ollama is not the selected AI provider.",
    })
    checks.append({
        "name": "semantic_matching",
        "status": "configured" if settings.semantic_matching_enabled else "disabled",
        "message": "Semantic matching is enabled in configuration; actual use depends on the embedding service/model being available." if settings.semantic_matching_enabled else "Semantic matching is disabled; rule-based scoring remains available.",
    })
    checks.append({
        "name": "brave_search",
        "status": "configured" if bool(settings.brave_search_api_key.strip()) else "not_configured",
        "message": "Broad web research has a key configured; provider availability is checked during search." if settings.brave_search_api_key.strip() else "Optional broad web research is not configured. Existing job feeds can still run.",
    })
    checks.append({
        "name": "adzuna",
        "status": "configured" if bool(settings.adzuna_app_id.strip() and settings.adzuna_app_key.strip()) else "not_configured",
        "message": "Adzuna credentials are configured." if settings.adzuna_app_id.strip() and settings.adzuna_app_key.strip() else "Optional Adzuna discovery is not configured.",
    })
    checks.append({
        "name": "jobspy",
        "status": "enabled" if settings.jobspy_enabled else "disabled",
        "message": "JobSpy is enabled by configuration; availability and each site result are reported during searches." if settings.jobspy_enabled else "Optional JobSpy multi-portal connector is disabled in backend configuration.",
    })
    return {"status": "ok" if overall_ok else "degraded", "version": APP_VERSION, "checks": checks}


@app.get("/api/jobs", response_model=list[JobRead], tags=["jobs"])
def list_jobs(limit: int = Query(default=50, ge=1, le=200), db: Session = Depends(get_db)):
    return db.scalars(select(Job).order_by(Job.created_at.desc()).limit(limit)).all()


@app.post("/api/jobs", response_model=JobRead, status_code=201, tags=["jobs"])
def create_job(payload: JobCreate, db: Session = Depends(get_db)):
    values = payload.model_dump()
    values["source_url"] = canonicalize_job_url(str(payload.source_url))
    job = Job(**values)
    db.add(job)
    try:
        db.commit()
        db.refresh(job)
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A job with this source URL already exists.")
    return job


def _source_api_payload(source: JobSource) -> dict[str, Any]:
    """Expose operational enablement, not just a persisted toggle.

    JobSpy requires both the source toggle and JOBSPY_ENABLED=true. If configuration
    disables it, showing enabled=true in the source list is misleading because the
    discovery pipeline will skip it. Keep the persisted user toggle intact so that
    explicitly enabling the connector in .env can restore it without losing settings.
    """
    effective_enabled = bool(source.enabled)
    if source.provider.lower() == "jobspy":
        effective_enabled = effective_enabled and bool(settings.jobspy_enabled)
    return {
        "id": source.id,
        "provider": source.provider,
        "company_name": source.company_name,
        "board_id": source.board_id,
        "enabled": effective_enabled,
        "created_at": source.created_at,
        "last_checked_at": source.last_checked_at,
    }


@app.get("/api/sources", response_model=list[JobSourceRead], tags=["sources"])
def list_sources(db: Session = Depends(get_db)):
    sources = db.scalars(select(JobSource).order_by(JobSource.created_at.desc())).all()
    return [_source_api_payload(source) for source in sources]


@app.post("/api/sources", response_model=JobSourceRead, status_code=201, tags=["sources"])
def create_source(payload: JobSourceCreate, db: Session = Depends(get_db)):
    board_id = payload.board_id.strip()
    provider = payload.provider.lower()
    existing = db.scalar(select(JobSource).where(JobSource.provider == provider, JobSource.board_id == board_id))
    if existing:
        raise HTTPException(status_code=409, detail="This provider and board identifier are already configured.")
    source = JobSource(provider=provider, company_name=payload.company_name.strip(), board_id=board_id, enabled=True)
    db.add(source)
    try:
        db.commit()
        db.refresh(source)
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="This job source already exists.")
    return _source_api_payload(source)


@app.patch("/api/sources/{source_id}", response_model=JobSourceRead, tags=["sources"])
def toggle_source(source_id: int, payload: JobSourceToggle, db: Session = Depends(get_db)):
    source = db.get(JobSource, source_id)
    if source is None:
        raise HTTPException(status_code=404, detail="Job source not found.")
    source.enabled = payload.enabled
    db.commit()
    db.refresh(source)
    return _source_api_payload(source)


@app.delete("/api/sources/{source_id}", status_code=204, tags=["sources"])
def delete_source(source_id: int, db: Session = Depends(get_db)):
    source = db.get(JobSource, source_id)
    if source is None:
        raise HTTPException(status_code=404, detail="Job source not found.")
    db.delete(source)
    db.commit()
    return None


async def _fetch_source(client: httpx.AsyncClient, source: JobSource, query: str = "") -> list[NormalizedPosting]:
    if source.provider == "greenhouse":
        return await fetch_greenhouse(client, source.board_id, source.company_name)
    if source.provider == "lever":
        return await fetch_lever(client, source.board_id, source.company_name)
    if source.provider == "arbeitnow":
        return await fetch_arbeitnow(client, source.board_id, source.company_name)
    if source.provider == "remoteok":
        return await fetch_remoteok(client, source.board_id, source.company_name)
    if source.provider == "jobicy":
        return await fetch_jobicy(client, source.board_id, source.company_name)
    if source.provider == "himalayas":
        return await fetch_himalayas(client, source.board_id, query, source.company_name)
    if source.provider == "adzuna":
        return await fetch_adzuna(client, source.board_id, query, source.company_name)
    if source.provider == "brave":
        # The full context-aware call is handled by discover_jobs, not this generic helper.
        raise ValueError("Brave Web Research requires search context; use the discovery endpoint.")
    if source.provider == "jobspy":
        raise ValueError("JobSpy requires query and location context; use the discovery endpoint.")
    raise ValueError(f"Unsupported source provider: {source.provider}")


def _to_discovery_item(
    posting: NormalizedPosting, age_hours: float | None = None, job_id: int | None = None,
    already_saved: bool = False, profile: CandidateProfile | None = None, query: str = "",
    qualification_reasons: list[str] | None = None,
) -> DiscoveryItem:
    details = score_posting(posting, profile, query) if profile else {}
    return DiscoveryItem(
        id=job_id, title=posting.title, company=posting.company, location=posting.location,
        work_mode=posting.work_mode, source_url=canonicalize_job_url(posting.source_url),
        description=posting.description, posted_at=posting.posted_at,
        verified_posting_time=posting.verified_posting_time, provider=posting.provider,
        source_company=posting.source_company, age_hours=age_hours, already_saved=already_saved,
        salary_range=details.get("salary_range", getattr(posting, "salary_range", "Not disclosed")),
        experience_required=details.get("experience_required", getattr(posting, "experience_required", "Not disclosed")),
        required_skills=details.get("required_skills", getattr(posting, "required_skills", [])),
        match_score=details.get("match_score", 0), rule_based_score=details.get("match_score", 0),
        semantic_fit_score=None, semantic_similarity=None, match_method=details.get("match_method", "rule_based"),
        matching_note="Explainable rule-based baseline; semantic ranking is optional.",
        score_breakdown=details.get("score_breakdown", {}), match_reasons=details.get("match_reasons", []),
        concerns=details.get("concerns", []), posting_time_evidence=details.get("posting_time_evidence", ""),
        verification_status=details.get("verification_status", "verified" if posting.verified_posting_time and posting.posted_at else "unverified"),
        qualification_reasons=qualification_reasons or [],
    )


async def _apply_semantic_scores_to_discovery_items(profile: CandidateProfile | None, items: list[DiscoveryItem]) -> None:
    """Enrich visible results with local embeddings while preserving the rule fallback."""
    if not items:
        return
    if profile is None:
        for item in items:
            item.match_method = "rule_based"
            item.matching_note = "No candidate profile is available; using the rule-based score."
        return
    scores = await semantic_scores_for_jobs(profile, items, [int(item.match_score or 0) for item in items])
    for item, score in zip(items, scores):
        item.rule_based_score = score.get("rule_based_score", item.match_score)
        item.semantic_fit_score = score.get("semantic_fit_score")
        item.semantic_similarity = score.get("semantic_similarity")
        item.match_score = score.get("match_score", item.match_score)
        item.match_method = score.get("match_method", "rule_based")
        item.matching_note = score.get("matching_note", "Using the rule-based score.")
        breakdown = dict(item.score_breakdown or {})
        if item.semantic_similarity is not None:
            breakdown["semantic_alignment"] = {
                "similarity": item.semantic_similarity,
                "fit_score": item.semantic_fit_score,
                "weight_percent": 30,
            }
        item.score_breakdown = breakdown
        if item.match_method == "hybrid_semantic":
            reasons = list(item.match_reasons or [])
            reasons.append("Semantic embedding alignment was included in the hybrid score; it is a ranking signal, not a hiring probability.")
            item.match_reasons = reasons[:8]


def _adzuna_countries_for_search(region: str, custom_location: str, configured_country: str) -> list[str]:
    """Choose Adzuna country endpoints for the selected search geography.

    Adzuna does not expose every country/region. For unsupported regions the other
    enabled sources still run, and the source status explains the limitation.
    """
    region = (region or "global").lower()
    grouped = {
        "global": ["in", "gb", "us", "ca", "au", "de", "fr", "nl", "sg", "pl", "za", "nz", "at", "be", "br", "ch", "es", "it", "mx"],
        "india": ["in"],
        "asia": ["in", "sg"],
        "europe": ["gb", "de", "fr", "nl", "pl", "at", "be", "ch", "es", "it"],
        "americas": ["us", "ca", "br", "mx"],
        "middle_east": [],
        "africa": ["za"],
        "oceania": ["au", "nz"],
    }
    if region == "custom":
        text_location = (custom_location or "").lower()
        patterns = {
            "india": ("india", "delhi", "mumbai", "bengaluru", "bangalore", "hyderabad", "pune", "gurugram", "noida"),
            "gb": ("united kingdom", "uk", "london", "england", "britain"),
            "us": ("united states", "usa", "u.s.", "new york", "california", "texas"),
            "ca": ("canada", "toronto", "vancouver"),
            "au": ("australia", "sydney", "melbourne"),
            "de": ("germany", "berlin", "munich"),
            "fr": ("france", "paris"),
            "nl": ("netherlands", "amsterdam"),
            "sg": ("singapore",),
            "za": ("south africa", "johannesburg", "cape town"),
        }
        for code, terms in patterns.items():
            if any(term in text_location for term in terms):
                return [code]
        return [configured_country.lower()] if configured_country.lower() in COUNTRY_CODES else []
    codes = grouped.get(region, [])
    return [code for code in codes if code in COUNTRY_CODES]


@app.post("/api/discovery/search", response_model=DiscoveryResponse, tags=["discovery"])
async def discover_jobs(payload: DiscoveryRequest, db: Session = Depends(get_db)):
    checked_at = datetime.now(timezone.utc)
    sources = db.scalars(select(JobSource).where(JobSource.enabled.is_(True)).order_by(JobSource.id)).all()
    # Copy ORM values before awaiting external network requests.
    configured_specs = [
        {"id": s.id, "provider": s.provider, "company_name": s.company_name, "board_id": s.board_id}
        for s in sources
    ]
    # Adzuna searches are expanded to region-specific country endpoints as a single
    # logical connector. The default setup in older databases may contain separate
    # India/UK/US rows, which would otherwise multiply every country request.
    source_specs = []
    skipped_adzuna_specs = []
    skipped_brave_specs = []
    skipped_jobspy_specs: list[tuple[dict[str, Any], str]] = []
    adzuna_primary_seen = False
    for spec in configured_specs:
        if spec["provider"] == "brave" and not settings.brave_search_api_key.strip():
            skipped_brave_specs.append(spec)
            continue
        if spec["provider"] == "jobspy":
            if not settings.jobspy_enabled:
                skipped_jobspy_specs.append((spec, "JobSpy is disabled by default. Install backend/requirements-jobspy.txt, set JOBSPY_ENABLED=true in backend/.env, restart, then enable this source."))
                continue
            if not jobspy_package_available():
                skipped_jobspy_specs.append((spec, "The optional python-jobspy package is not installed. Run `pip install -r requirements-jobspy.txt` in the backend environment, then restart."))
                continue
        if spec["provider"] == "adzuna":
            if adzuna_primary_seen:
                skipped_adzuna_specs.append(spec)
                continue
            adzuna_primary_seen = True
        source_specs.append(spec)
    profile = db.get(CandidateProfile, 1)
    query = payload.query.strip() or (profile.target_titles if profile and profile.target_titles else "")
    if not source_specs:
        skipped_statuses = [SourceCheckStatus(
            source_id=spec["id"], provider="brave", company_name=spec["company_name"], status="skipped",
            fetched_count=0, message="Optional broad-web search is not configured. Add BRAVE_SEARCH_API_KEY to backend/.env and restart the backend.",
        ) for spec in skipped_brave_specs]
        skipped_statuses.extend(SourceCheckStatus(
            source_id=spec["id"], provider="jobspy", company_name=spec["company_name"], status="skipped",
            fetched_count=0, message=reason,
        ) for spec, reason in skipped_jobspy_specs)
        return DiscoveryResponse(
            query=query, hours=payload.hours, checked_at=checked_at, sources_checked=0,
            fresh_count=0, unverified_count=0, fresh_jobs=[], unverified_jobs=[],
            source_statuses=skipped_statuses, region=payload.region, custom_location=payload.custom_location,
            work_modes=payload.work_modes, employment_types=payload.employment_types,
            web_search_configured=bool(settings.brave_search_api_key.strip()),
            web_search_notes=["No active public job sources are configured."] + ([] if settings.brave_search_api_key.strip() else ["Configure BRAVE_SEARCH_API_KEY to enable broad web research."]),
        )

    timeout = httpx.Timeout(15.0, connect=7.0)
    raw: list[NormalizedPosting] = []
    statuses: list[SourceCheckStatus] = []
    semaphore = asyncio.Semaphore(6)

    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True, headers={"Accept": "application/json", "User-Agent": "AI-Job-Hunter/0.2 (open-source job discovery)"}) as client:
        async def check_one(spec):
            async with semaphore:
                try:
                    # Rehydrate the small source object needed by the connector.
                    class SourceSpec:
                        pass
                    source = SourceSpec()
                    source.provider = spec["provider"]
                    source.board_id = spec["board_id"]
                    source.company_name = spec["company_name"]
                    if source.provider == "brave":
                        postings, diagnostics = await fetch_brave_web_research(
                            client=client, api_key=settings.brave_search_api_key, query=query,
                            region=payload.region, custom_location=payload.custom_location, hours=payload.hours,
                        )
                        spec["web_research_diagnostics"] = diagnostics
                    elif source.provider == "himalayas":
                        postings = await fetch_himalayas(client, source.board_id, query, source.company_name)
                    elif source.provider == "adzuna":
                        country_codes = _adzuna_countries_for_search(payload.region, payload.custom_location, source.board_id)
                        if not country_codes:
                            raise ValueError(
                                f"Adzuna has no configured country endpoint for {payload.region.replace('_', ' ')}; "
                                "other enabled sources are still searched."
                            )
                        country_sem = asyncio.Semaphore(4)
                        async def fetch_country(code):
                            async with country_sem:
                                return await fetch_adzuna(client, code, query, f"{source.company_name} ({code.upper()})")
                        results = await asyncio.gather(*(fetch_country(code) for code in country_codes), return_exceptions=True)
                        postings = []
                        failures = []
                        for code, result in zip(country_codes, results):
                            if isinstance(result, Exception):
                                failures.append(f"{code.upper()}: {type(result).__name__}")
                            else:
                                postings.extend(result)
                        if not postings and failures:
                            raise RuntimeError("Adzuna requests failed for all selected markets (" + ", ".join(failures[:5]) + "). Check API credentials, connectivity, or rate limits.")
                    elif source.provider == "jobspy":
                        postings, diagnostics = await fetch_jobspy_postings(
                            query=query,
                            region=payload.region,
                            custom_location=payload.custom_location,
                            hours=payload.hours,
                            sites=[site.strip().lower() for site in settings.jobspy_sites.split(",") if site.strip()],
                            max_searches=settings.jobspy_max_searches,
                            results_wanted=settings.jobspy_results_wanted,
                        )
                        spec["jobspy_diagnostics"] = diagnostics
                    else:
                        # Keep this two-argument call for straightforward connector mocking in tests.
                        postings = await _fetch_source(client, source)
                    return spec, postings, None
                except httpx.HTTPStatusError as exc:
                    # A bad board token is a configuration problem, not a reason to
                    # fail the entire search. Pause a Lever board after 404 so a stale
                    # placeholder does not generate the same traceback on every run.
                    status_code = exc.response.status_code if exc.response is not None else None
                    if spec["provider"] == "lever" and status_code == 404:
                        spec["auto_pause_source"] = True
                        reason = (
                            f"Lever returned HTTP 404 for board ID '{spec['board_id']}'. "
                            "This usually means the site slug is incorrect, is a test placeholder, "
                            "or the company board no longer exists. The source has been paused to "
                            "prevent repeated failures; remove it and add a valid Lever site slug."
                        )
                        logger.warning("Pausing invalid Lever job source (company=%s, board=%s, status=%s)",
                                       spec["company_name"], spec["board_id"], status_code)
                    else:
                        reason = f"HTTP {status_code} from source: {str(exc).strip()}".strip()
                        logger.warning("Job source HTTP error (provider=%s, board=%s, status=%s): %s",
                                       spec["provider"], spec["board_id"], status_code, str(exc)[:300])
                    return spec, [], (reason or "Unexpected HTTP error from connector")[:420]
                except Exception as exc:
                    # A malformed response from one source should not fail the whole search.
                    logger.exception("Job source check failed (provider=%s, board=%s)", spec["provider"], spec["board_id"])
                    reason = f"{type(exc).__name__}: {str(exc).strip()}".strip()
                    return spec, [], (reason or "Unexpected connector error")[:240]

        outcomes = await asyncio.gather(*(check_one(spec) for spec in source_specs))

    for spec, postings, error in outcomes:
        if error:
            statuses.append(SourceCheckStatus(source_id=spec["id"], provider=spec["provider"], company_name=spec["company_name"], status="error", fetched_count=0, message=error))
        else:
            raw.extend(postings)
            if spec["provider"] == "jobspy":
                diag = spec.get("jobspy_diagnostics", {})
                search_succeeded = int(diag.get("searches_succeeded", 0))
                errors = diag.get("errors", [])
                if not postings and search_succeeded == 0 and errors:
                    source_status = "error"
                    message = "JobSpy could not complete any search. " + " | ".join(errors[:3])
                elif not postings:
                    source_status = "empty"
                    message = (f"JobSpy completed {search_succeeded}/{diag.get('searches_attempted', 0)} search variants and returned 0 postings. "
                               "This is different from a source error.")
                else:
                    source_status = "success"
                    message = (f"JobSpy fetched {len(postings)} postings across {diag.get('searches_succeeded', 0)} successful role/location searches; "
                               f"sites attempted: {', '.join(diag.get('sites_attempted', [])) or 'none'}. "
                               "Publication times without time-of-day remain unverified.")
                if errors:
                    message += " Partial source errors: " + " | ".join(errors[:3])
                if diag.get("warnings"):
                    message += " Warnings: " + " | ".join(diag["warnings"][:2])
            elif not postings:
                source_status = "empty"
                message = "Request completed successfully, but this source returned 0 postings for this run."
            elif spec["provider"] == "brave":
                source_status = "success"
                diag = spec.get("web_research_diagnostics", {})
                message = (f"Web research: {diag.get('queries_succeeded', 0)}/{diag.get('queries_planned', 0)} queries succeeded; "
                           f"{diag.get('unique_candidates_checked', len(postings))} unique candidate pages checked; "
                           f"{diag.get('original_pages_verified', 0)} had source-level posting-time evidence.")
            else:
                source_status = "success"
                message = f"Fetched {len(postings)} published postings"
            statuses.append(SourceCheckStatus(source_id=spec["id"], provider=spec["provider"], company_name=spec["company_name"], status=source_status, fetched_count=len(postings), message=message))
        source = db.get(JobSource, spec["id"])
        if source:
            source.last_checked_at = checked_at
            if spec.get("auto_pause_source"):
                source.enabled = False
    for spec in skipped_adzuna_specs:
        statuses.append(SourceCheckStatus(
            source_id=spec["id"], provider=spec["provider"], company_name=spec["company_name"],
            status="skipped", fetched_count=0,
            message="Covered by the region-aware Adzuna connector; duplicate country configuration was not fetched separately.",
        ))
        source = db.get(JobSource, spec["id"])
        if source:
            source.last_checked_at = checked_at
    for spec in skipped_brave_specs:
        statuses.append(SourceCheckStatus(
            source_id=spec["id"], provider="brave", company_name=spec["company_name"], status="skipped",
            fetched_count=0, message="Not configured: add BRAVE_SEARCH_API_KEY to backend/.env and restart. This optional connector searches the wider web; other sources continue without it.",
        ))
    for spec, reason in skipped_jobspy_specs:
        statuses.append(SourceCheckStatus(
            source_id=spec["id"], provider="jobspy", company_name=spec["company_name"],
            status="skipped", fetched_count=0, message=reason,
        ))
    db.commit()

    web_diag = next((spec.get("web_research_diagnostics", {}) for spec in source_specs if spec["provider"] == "brave"), {})
    web_search_notes = []
    for error in web_diag.get("errors", []):
        web_search_notes.append(f"Web search request: {error}")
    if settings.brave_search_api_key.strip() and web_diag:
        web_search_notes.append(
            f"Web research ran {web_diag.get('queries_succeeded', 0)}/{web_diag.get('queries_planned', 0)} queries and checked {web_diag.get('unique_candidates_checked', 0)} original pages; {web_diag.get('original_pages_verified', 0)} had source-level posting time evidence."
        )
        portal_breakdown = web_diag.get("portal_breakdown", [])
        if portal_breakdown:
            portal_summary = "; ".join(
                f"{item.get('portal', 'Unknown site')}: {item.get('results_seen', 0)} result leads"
                for item in portal_breakdown[:12]
            )
            web_search_notes.append("Indexed portal results by source domain — " + portal_summary + ".")
        query_breakdown = web_diag.get("query_breakdown", [])
        if query_breakdown:
            query_summary = "; ".join(
                f"{item.get('scope', 'Search')}: {item.get('succeeded', 0)}/{item.get('planned', 0)} queries succeeded, {item.get('results_seen', 0)} indexed results"
                for item in query_breakdown[:12]
            )
            web_search_notes.append("Search scope coverage — " + query_summary + ".")

    phrases = query_phrases(query)
    raw_postings_count = len(raw)
    unique_raw, duplicate_pairs = deduplicate_postings(raw)
    exclusions: list[dict[str, Any]] = []
    exclusion_counts: Counter[str] = Counter()

    def record_exclusion(posting: NormalizedPosting, reason_code: str, reason: str):
        exclusion_counts[reason_code] += 1
        exclusions.append({
            "title": (posting.title or "Untitled listing")[:250],
            "company": (posting.company or "Unknown company")[:250],
            "location": (posting.location or "Location not specified")[:250],
            "source_url": canonicalize_job_url(posting.source_url or ""),
            "provider": posting.provider or "unknown",
            "reason_code": reason_code,
            "reason": reason[:500],
            "posted_at": posting.posted_at,
            "posting_time_evidence": posting.posting_time_evidence or ("Source timestamp provided" if posting.verified_posting_time else "Original publication time not verified"),
        })

    for duplicate, kept in duplicate_pairs:
        record_exclusion(
            duplicate, "duplicate_listing",
            f"Duplicate listing of '{kept.title}' at {kept.company}; the other copy was retained from {kept.source_company or kept.provider}.",
        )

    query_matched_all: list[NormalizedPosting] = []
    query_reasons: dict[int, str] = {}
    for posting in unique_raw:
        if not posting.title or not posting.source_url:
            record_exclusion(posting, "invalid_listing", "Listing is missing a title or direct source URL, so it cannot be reliably reviewed or applied to.")
            continue
        matched, why = query_match_reason(posting, query, phrases)
        if not matched:
            record_exclusion(posting, "role_keyword_mismatch", why)
            continue
        query_matched_all.append(posting)
        query_reasons[id(posting)] = why

    region_matched: list[NormalizedPosting] = []
    for posting in query_matched_all:
        if not location_matches_region(posting, payload.region, payload.custom_location):
            label = payload.custom_location if payload.region == "custom" else payload.region.replace("_", " ").title()
            record_exclusion(posting, "location_mismatch", f"Listing location '{posting.location or 'not specified'}' did not match selected search region '{label}'.")
            continue
        region_matched.append(posting)

    mode_matched: list[NormalizedPosting] = []
    selected_modes = set(payload.work_modes)
    for posting in region_matched:
        detected = classify_work_mode(posting)
        if detected not in selected_modes:
            selected = ", ".join(payload.work_modes) or "none"
            record_exclusion(posting, "work_mode_mismatch", f"Detected work arrangement '{detected}' was not selected. Selected values: {selected}.")
            continue
        mode_matched.append(posting)

    employment_matched: list[NormalizedPosting] = []
    selected_employment = set(payload.employment_types)
    for posting in mode_matched:
        detected = classify_employment_type(posting)
        if detected not in selected_employment:
            selected = ", ".join(payload.employment_types) or "none"
            record_exclusion(posting, "employment_type_mismatch", f"Detected employment type '{detected}' was not selected. Selected values: {selected}.")
            continue
        employment_matched.append(posting)

    cutoff = checked_at - timedelta(hours=payload.hours)
    time_window_excluded = 0
    timestamp_unverified_candidates = 0
    verified_fresh_candidates = 0
    for posting in employment_matched:
        if posting.verified_posting_time and posting.posted_at:
            published = posting.posted_at
            if published.tzinfo is None:
                published = published.replace(tzinfo=timezone.utc)
            published = published.astimezone(timezone.utc)
            if published > checked_at:
                time_window_excluded += 1
                record_exclusion(posting, "future_or_invalid_timestamp", f"Source publication timestamp ({published.isoformat()}) is later than the search time; it was not treated as fresh.")
            elif published < cutoff:
                time_window_excluded += 1
                age_days = (checked_at - published).total_seconds() / 86400
                record_exclusion(posting, "outside_posting_window", f"Source publication timestamp is about {age_days:.1f} days old, outside the selected {payload.hours}-hour window.")
            else:
                verified_fresh_candidates += 1
        else:
            timestamp_unverified_candidates += 1
            reason = "Original posting date could not be verified, so this lead is excluded from Fresh Matches."
            if payload.include_unverified:
                reason += " It is included separately in Unverified leads."
            else:
                reason += " Enable 'Also show listings where the source does not provide an original posting timestamp' to inspect it."
            record_exclusion(posting, "posting_time_unverified", reason)

    location_filtered_out_count = exclusion_counts.get("location_mismatch", 0)
    work_mode_filtered_out_count = exclusion_counts.get("work_mode_mismatch", 0)
    employment_type_filtered_out_count = exclusion_counts.get("employment_type_mismatch", 0)
    query_filtered_out_count = len(unique_raw) - len(query_matched_all)

    # Only in-scope records are eligible to be shown/saved. Unknown timestamps stay in a
    # separate bucket; they are never promoted to verified fresh matches.
    fresh, unverified = filter_and_deduplicate(employment_matched, query="", hours=payload.hours, now=checked_at)
    saved_jobs = db.scalars(select(Job)).all()
    saved_by_url = {canonicalize_job_url(job.source_url): job for job in saved_jobs}
    saved_by_fingerprint = {(" ".join(job.title.lower().split()), " ".join(job.company.lower().split()), " ".join(job.location.lower().split())): job for job in saved_jobs}

    def qualification_reasons(posting: NormalizedPosting, age_hours: float | None) -> list[str]:
        reasons = [query_reasons.get(id(posting), "Matched the selected role query.")]
        region_label = payload.custom_location if payload.region == "custom" else payload.region.replace("_", " ")
        reasons.append(f"Location scope: matched selected region '{region_label}'.")
        reasons.append(f"Work arrangement '{classify_work_mode(posting)}' is included in your filters.")
        reasons.append(f"Employment type '{classify_employment_type(posting)}' is included in your filters.")
        if age_hours is not None:
            reasons.append(f"Freshness: the source provides a timestamp about {age_hours:.1f} hours before this search.")
        reasons.append(f"Source: {posting.source_company or posting.provider}.")
        return reasons

    fresh_items: list[DiscoveryItem] = []
    seen_fingerprints = set()
    for posting, age_hours in fresh:
        canonical_url = canonicalize_job_url(posting.source_url)
        semantic_key = (" ".join(posting.title.lower().split()), " ".join(posting.company.lower().split()), " ".join(posting.location.lower().split()))
        if semantic_key in seen_fingerprints:
            record_exclusion(posting, "duplicate_listing", "Duplicate title/company/location combination; another copy was kept.")
            continue
        seen_fingerprints.add(semantic_key)
        existing = saved_by_url.get(canonical_url) or saved_by_fingerprint.get(semantic_key)
        qualification = qualification_reasons(posting, age_hours)
        if existing:
            fresh_items.append(_to_discovery_item(posting, age_hours, existing.id, True, profile=profile, query=query, qualification_reasons=qualification))
            continue
        job = Job(
            title=posting.title[:250], company=posting.company[:250], location=posting.location[:250],
            work_mode=posting.work_mode[:50], source_url=canonical_url,
            description=posting.description[:20000], posted_at=posting.posted_at,
            verified_posting_time=True,
        )
        db.add(job)
        try:
            db.flush()
            saved_by_url[canonical_url] = job
            saved_by_fingerprint[semantic_key] = job
            fresh_items.append(_to_discovery_item(posting, age_hours, job.id, False, profile=profile, query=query, qualification_reasons=qualification))
        except IntegrityError:
            db.rollback()
            existing = db.scalar(select(Job).where(Job.source_url == canonical_url))
            if existing:
                fresh_items.append(_to_discovery_item(posting, age_hours, existing.id, True, profile=profile, query=query, qualification_reasons=qualification))
    db.commit()

    unknown_items: list[DiscoveryItem] = []
    if payload.include_unverified:
        for posting in unverified:
            unknown_items.append(_to_discovery_item(
                posting, profile=profile, query=query,
                qualification_reasons=[query_reasons.get(id(posting), "Matched the selected role query."),
                    "Location, work arrangement, and employment type matched the selected filters.",
                    "Not included in Fresh Matches because the original posting time could not be verified."],
            ))

    await _apply_semantic_scores_to_discovery_items(profile, fresh_items + unknown_items)

    # Funnel counts explain how the raw feed becomes the final job list.
    unique_count = len(unique_raw)
    funnel = [
        {"stage": "Fetched", "description": "All records returned by successful configured sources before deduplication.", "input_count": raw_postings_count, "passed_count": unique_count, "excluded_count": raw_postings_count - unique_count},
        {"stage": "Role relevance", "description": "Title variants and selected role keywords were checked against each listing.", "input_count": unique_count, "passed_count": len(query_matched_all), "excluded_count": query_filtered_out_count},
        {"stage": "Location", "description": "Listings were checked against the selected region/custom location.", "input_count": len(query_matched_all), "passed_count": len(region_matched), "excluded_count": len(query_matched_all) - len(region_matched)},
        {"stage": "Work arrangement", "description": "Only selected remote/hybrid/on-site/unknown values were kept.", "input_count": len(region_matched), "passed_count": len(mode_matched), "excluded_count": len(region_matched) - len(mode_matched)},
        {"stage": "Employment type", "description": "Only selected full-time/part-time/contract/internship/temporary/unknown values were kept.", "input_count": len(mode_matched), "passed_count": len(employment_matched), "excluded_count": len(mode_matched) - len(employment_matched)},
        {"stage": "Posting window", "description": "Only source-timestamped listings within the selected time window become Fresh Matches; unknown dates are shown separately if enabled.", "input_count": len(employment_matched), "passed_count": len(fresh_items), "excluded_count": max(0, len(employment_matched) - len(fresh_items))},
    ]
    # Keep a useful audit view without allowing an unexpectedly large source to create an
    # unbounded API response; totals and reason counts always cover every excluded record.
    exclusion_cap = 500
    exposed_exclusions = exclusions[:exclusion_cap]

    return DiscoveryResponse(
        query=query, hours=payload.hours, checked_at=checked_at, sources_checked=len(source_specs),
        fresh_count=len(fresh_items), unverified_count=len(unknown_items),
        fresh_jobs=fresh_items, unverified_jobs=unknown_items, source_statuses=statuses,
        region=payload.region, custom_location=payload.custom_location,
        work_modes=payload.work_modes, employment_types=payload.employment_types,
        raw_postings_count=raw_postings_count, unique_postings_count=unique_count,
        duplicate_removed_count=raw_postings_count - unique_count,
        query_filtered_out_count=query_filtered_out_count,
        query_matched_count=len(query_matched_all),
        location_matched_count=len(region_matched), work_mode_matched_count=len(mode_matched),
        employment_type_matched_count=len(employment_matched),
        location_filtered_out_count=location_filtered_out_count,
        work_mode_filtered_out_count=work_mode_filtered_out_count,
        employment_type_filtered_out_count=employment_type_filtered_out_count,
        time_window_excluded_count=time_window_excluded,
        timestamp_unverified_candidates_count=timestamp_unverified_candidates,
        excluded_listing_count=len(exclusions), excluded_listings_truncated=len(exclusions) > exclusion_cap,
        exclusion_reason_counts=dict(exclusion_counts), excluded_listings=exposed_exclusions,
        filter_funnel=funnel,
        web_search_configured=bool(settings.brave_search_api_key.strip()),
        web_search_queries=web_diag.get("queries_succeeded", 0),
        web_search_original_pages_verified=web_diag.get("original_pages_verified", 0),
        web_search_notes=web_search_notes,
    )


@app.get("/api/profile", response_model=ProfileRead, tags=["profile"])
def get_profile(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        profile = CandidateProfile(id=1)
        db.add(profile)
        db.commit()
        db.refresh(profile)
    return profile


@app.put("/api/profile", response_model=ProfileRead, tags=["profile"])
def update_profile(payload: ProfileUpdate, db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        profile = CandidateProfile(id=1)
        db.add(profile)
    for key, value in payload.model_dump().items():
        setattr(profile, key, value)
    db.commit()
    db.refresh(profile)
    return profile


@app.post("/api/profile/resume", response_model=ResumeUploadResponse, tags=["profile"])
async def upload_profile_resume(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upload a private resume file (2 MB max) and extract selectable text where supported."""
    original_name = re.split(r"[/\\]", file.filename or "resume")[-1][:255]
    suffix = Path(original_name).suffix.lower()
    if suffix not in SUPPORTED_RESUME_SUFFIXES:
        raise HTTPException(status_code=415, detail="Supported formats: PDF, DOC, DOCX, and RTF.")
    data = await file.read(MAX_RESUME_BYTES + 1)
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume file must be 2 MB or smaller.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF.")

    try:
        extracted_text, parse_status = _extract_resume_text(data, suffix)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this resume. It may be encrypted, damaged, or in an unsupported variant.") from exc
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        profile = CandidateProfile(id=1)
        db.add(profile)
        db.flush()

    storage_key = f"{uuid4().hex}{suffix}"
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    destination = RESUME_DIR / storage_key
    try:
        destination.write_bytes(data)
    except OSError as exc:
        raise HTTPException(status_code=500, detail="Could not store the resume locally.") from exc

    previous_path = _safe_resume_path(profile.resume_storage_key)
    profile.resume_storage_key = storage_key
    profile.resume_original_filename = original_name
    profile.resume_uploaded_at = datetime.now(timezone.utc)
    profile.resume_size_bytes = len(data)
    profile.resume_parse_status = parse_status
    profile.resume_text = extracted_text
    db.commit()
    if previous_path and previous_path.exists() and previous_path != destination:
        previous_path.unlink(missing_ok=True)

    return ResumeUploadResponse(
        filename=original_name,
        size_bytes=len(data),
        uploaded_at=profile.resume_uploaded_at,
        parse_status=parse_status,
        extracted_characters=len(extracted_text),
        extracted_text_preview=extracted_text[:3000],
    )


@app.get("/api/profile/resume/text", response_model=ResumeTextResponse, tags=["profile"])
def get_profile_resume_text(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    if profile is None or not profile.resume_storage_key:
        raise HTTPException(status_code=404, detail="No resume has been uploaded yet.")
    return ResumeTextResponse(
        filename=profile.resume_original_filename or "resume",
        parse_status=profile.resume_parse_status,
        text=profile.resume_text or "",
    )


@app.get("/api/profile/resume/download", tags=["profile"])
def download_profile_resume(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    path = _safe_resume_path(profile.resume_storage_key if profile else None)
    if not path or not path.is_file():
        raise HTTPException(status_code=404, detail="No resume file is available.")
    return FileResponse(
        path,
        filename=profile.resume_original_filename or "resume",
        media_type="application/octet-stream",
    )


@app.delete("/api/profile/resume", status_code=204, tags=["profile"])
def delete_profile_resume(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    if profile is None or not profile.resume_storage_key:
        raise HTTPException(status_code=404, detail="No resume has been uploaded yet.")
    path = _safe_resume_path(profile.resume_storage_key)
    if path and path.exists():
        path.unlink(missing_ok=True)
    profile.resume_storage_key = None
    profile.resume_original_filename = None
    profile.resume_uploaded_at = None
    profile.resume_size_bytes = None
    profile.resume_parse_status = "not_uploaded"
    profile.resume_text = ""
    db.commit()
    return None


@app.post("/api/profile/photo", response_model=PhotoUploadResponse, tags=["profile"])
async def upload_profile_photo(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Store an optional profile avatar locally. Images are type-checked by extension and signature."""
    original_name = re.split(r"[/\\\\]", file.filename or "profile-photo")[-1][:255]
    suffix = Path(original_name).suffix.lower()
    allowed = {
        ".png": ("image/png", "PNG"),
        ".jpg": ("image/jpeg", "JPEG"),
        ".jpeg": ("image/jpeg", "JPEG"),
        ".webp": ("image/webp", "WEBP"),
    }
    if suffix not in allowed:
        raise HTTPException(status_code=415, detail="Profile photo must be PNG, JPG, JPEG, or WEBP.")
    data = await file.read(MAX_RESUME_BYTES + 1)
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Profile photo must be 2 MB or smaller.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded image is empty.")
    content_type, expected_format = allowed[suffix]
    try:
        from PIL import Image, UnidentifiedImageError
        with Image.open(io.BytesIO(data)) as image:
            if image.format != expected_format:
                raise HTTPException(status_code=400, detail="The file content does not match the image extension.")
            width, height = image.size
            if width < 1 or height < 1 or width > 5000 or height > 5000 or width * height > 20_000_000:
                raise HTTPException(status_code=413, detail="Profile photo dimensions exceed the 5000 × 5000 safety limit.")
            image.verify()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid supported image.") from exc

    profile = db.get(CandidateProfile, 1)
    if profile is None:
        profile = CandidateProfile(id=1)
        db.add(profile)
        db.flush()
    photo_dir = RESUME_DIR.parent / "photos"
    photo_dir.mkdir(parents=True, exist_ok=True)
    key = f"{uuid4().hex}{suffix}"
    (photo_dir / key).write_bytes(data)
    previous = _safe_photo_path(profile.photo_storage_key)
    uploaded_at = datetime.now(timezone.utc)
    profile.photo_storage_key = key
    profile.photo_original_filename = original_name
    profile.photo_content_type = content_type
    profile.photo_size_bytes = len(data)
    profile.photo_uploaded_at = uploaded_at
    db.commit()
    if previous and previous.exists():
        previous.unlink(missing_ok=True)
    return PhotoUploadResponse(filename=original_name, size_bytes=len(data), content_type=content_type, uploaded_at=uploaded_at)


@app.get("/api/profile/photo", tags=["profile"])
def get_profile_photo(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    path = _safe_photo_path(profile.photo_storage_key if profile else None)
    if not path or not path.is_file():
        raise HTTPException(status_code=404, detail="No profile photo has been uploaded.")
    return FileResponse(path, media_type=profile.photo_content_type or "application/octet-stream", filename=profile.photo_original_filename or "profile-photo")


@app.delete("/api/profile/photo", status_code=204, tags=["profile"])
def delete_profile_photo(db: Session = Depends(get_db)):
    profile = db.get(CandidateProfile, 1)
    if profile is None or not profile.photo_storage_key:
        raise HTTPException(status_code=404, detail="No profile photo has been uploaded.")
    path = _safe_photo_path(profile.photo_storage_key)
    if path and path.exists():
        path.unlink(missing_ok=True)
    profile.photo_storage_key = None
    profile.photo_original_filename = None
    profile.photo_content_type = None
    profile.photo_size_bytes = None
    profile.photo_uploaded_at = None
    db.commit()
    return None


@app.post("/api/profile/resume/suggestions", tags=["profile"])
def suggest_profile_from_resume(db: Session = Depends(get_db)):
    """Provide editable, heuristic suggestions from extracted resume text; nothing is auto-saved."""
    profile = db.get(CandidateProfile, 1)
    if profile is None or not (profile.resume_text or "").strip():
        raise HTTPException(status_code=400, detail="Upload a resume with selectable text before generating suggestions.")
    text_content = profile.resume_text.strip()
    lines = [re.sub(r"\s+", " ", line).strip(" •-\t") for line in text_content.splitlines() if line.strip()]
    headline_terms = ("recruiter", "talent acquisition", "human resources", "hr executive", "data scientist", "data analyst", "machine learning", "software engineer", "product manager")
    headline = next((line for line in lines[:35] if len(line) <= 220 and any(term in line.lower() for term in headline_terms)), "")
    paragraphs = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n{2,}", text_content) if p.strip()]
    summary = next((p for p in paragraphs if 80 <= len(p) <= 1200 and any(term in p.lower() for term in ("experience", "responsible", "recruitment", "professional summary", "profile summary", "worked on"))), "")
    skill_dictionary = [
        "IT Recruitment", "Non-IT Recruitment", "Talent Acquisition", "End-to-End Recruitment", "Boolean Search",
        "LinkedIn Recruiter", "Stakeholder Management", "Candidate Sourcing", "Talent Mapping", "Interview Coordination",
        "Generative AI", "Data Science", "Machine Learning", "Python", "SQL", "Power BI", "Pandas", "Excel",
        "SAP", "Salesforce", "Technical Recruitment", "International Recruitment", "Campus Recruitment",
    ]
    lowered = text_content.lower()
    skills = [skill for skill in skill_dictionary if skill.lower() in lowered]
    return {
        "headline": headline,
        "profile_summary": summary,
        "skills": ", ".join(skills),
        "matched_skill_terms": skills,
        "note": "Suggestions are heuristic extracts from your resume text. Review every field and edit it before saving; no profile changes have been made.",
    }


def _profile_keyword_context(profile: CandidateProfile) -> tuple[set[str], list[str], str]:
    fields = [profile.target_titles, profile.skills, profile.resume_headline, profile.profile_summary, profile.resume_text]
    try:
        career = __import__("json").loads(profile.career_profile_json or "{}")
        fields.extend(str(value) for value in career.values() if value is not None)
    except (ValueError, AttributeError):
        pass
    tokens = set(re.findall(r"[a-zA-Z][a-zA-Z0-9+#.-]{1,}", " ".join(fields).lower()))
    skills = [part.strip().lower() for part in re.split(r"[,;\n]", profile.skills or "") if len(part.strip()) > 1]
    return tokens, skills, " ".join(fields).lower()


def _rule_job_match_rows(profile: CandidateProfile, jobs: list[Job]) -> list[dict[str, Any]]:
    """Fast deterministic ranking used as the always-available matching baseline."""
    tokens, skills, _ = _profile_keyword_context(profile)
    rows: list[dict[str, Any]] = []
    target_phrases = [x.strip().lower() for x in (profile.target_titles or "").split(",") if x.strip()]
    try:
        career = json.loads(profile.career_profile_json or "{}")
        preferred_locations = [x.strip().lower() for x in str(career.get("preferred_locations", profile.preferred_locations)).split(",") if x.strip()]
    except (ValueError, AttributeError):
        preferred_locations = [x.strip().lower() for x in (profile.preferred_locations or "").split(",") if x.strip()]
    for job in jobs:
        title = job.title.lower()
        description = (job.description or "").lower()
        body = f"{title} {description}"
        matched = sorted({skill for skill in skills if skill in body})
        matched_tokens = {token for token in tokens if len(token) >= 3 and re.search(rf"(?<![a-z0-9]){re.escape(token)}(?![a-z0-9])", body)}
        denominator = max(1, min(len(tokens), 35))
        score = min(78, round(78 * min(1.0, len(matched_tokens) / denominator)))
        reasons = []
        if any(phrase in title for phrase in target_phrases):
            score = min(100, score + 15)
            reasons.append("Job title matches one of your target roles.")
        if matched:
            score = min(100, score + min(20, len(matched) * 4))
            reasons.append(f"Matched skills: {', '.join(matched[:8])}.")
        if preferred_locations and any(loc in (job.location or "").lower() for loc in preferred_locations):
            score = min(100, score + 7)
            reasons.append("Location matches your preferences.")
        if not reasons:
            reasons.append("Rule-based keyword baseline; review the job description manually.")
        rows.append({"job": job, "rule_score": max(0, min(100, score)), "matched": matched[:12], "reasons": reasons})
    return rows


def _job_match_records(rows: list[dict[str, Any]], semantic: list[dict[str, Any]] | None = None) -> list[JobMatchRead]:
    records = []
    for index, row in enumerate(rows):
        score_data = semantic[index] if semantic and index < len(semantic) else {
            "rule_based_score": row["rule_score"],
            "semantic_fit_score": None,
            "semantic_similarity": None,
            "match_score": row["rule_score"],
            "match_method": "rule_based",
            "matching_note": "Semantic re-ranking was not run; using the rule-based score.",
        }
        reasons = list(row["reasons"])
        if score_data.get("match_method") == "hybrid_semantic":
            reasons.append("Rank includes local embedding similarity; this is a ranking signal, not a hiring probability.")
        records.append(JobMatchRead(
            job=row["job"], match_score=score_data.get("match_score", row["rule_score"]),
            matched_keywords=row["matched"], reasons=reasons,
            rule_based_score=score_data.get("rule_based_score", row["rule_score"]),
            semantic_fit_score=score_data.get("semantic_fit_score"),
            semantic_similarity=score_data.get("semantic_similarity"),
            match_method=score_data.get("match_method", "rule_based"),
            matching_note=score_data.get("matching_note", "Using a rule-based score."),
            score_breakdown={},
        ))
    records.sort(key=lambda item: item.match_score, reverse=True)
    return records


def _get_rule_job_matches(limit: int, db: Session) -> list[JobMatchRead]:
    """Synchronous baseline for existing sync routes that create application records."""
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        return []
    jobs = db.scalars(select(Job).order_by(Job.created_at.desc()).limit(limit)).all()
    return _job_match_records(_rule_job_match_rows(profile, jobs))


@app.get("/api/jobs/matches", response_model=list[JobMatchRead], tags=["matching"])
async def get_job_matches(limit: int = Query(default=100, ge=1, le=200), db: Session = Depends(get_db)):
    """Return explainable match scores with optional local semantic re-ranking."""
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        return []
    jobs = db.scalars(select(Job).order_by(Job.created_at.desc()).limit(limit)).all()
    rows = _rule_job_match_rows(profile, jobs)
    semantic = await semantic_scores_for_jobs(profile, [row["job"] for row in rows], [row["rule_score"] for row in rows])
    return _job_match_records(rows, semantic)


@app.post("/api/jobs/{job_id}/analysis", response_model=JobAnalysisResponse, tags=["matching"])
async def analyze_saved_job(job_id: int, db: Session = Depends(get_db)):
    """Return a resume-grounded job-fit breakdown and suggested screening questions.

    The AI portion is advisory and falls back to deterministic evidence when the
    configured provider is unavailable. It never submits an application.
    """
    job = db.get(Job, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Saved job not found. Run discovery or save the job before analyzing it.")
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        raise HTTPException(status_code=400, detail="Complete your candidate profile before analyzing a job.")
    result = await analyze_job_fit(profile, job)
    return JobAnalysisResponse(**result)


# ---------------------- Application history and agent workflow ----------------------


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _normalise_platform(source_url: str) -> str:
    return platform_from_url(source_url)


def _application_for_job(job: Job, score: int, matched_keywords: list[str], profile: CandidateProfile) -> Application:
    app = Application(
        job_id=job.id,
        title=job.title,
        company=job.company,
        location=job.location or "Not specified",
        work_mode=job.work_mode or "Unknown",
        source_url=canonicalize_job_url(job.source_url),
        platform=_normalise_platform(job.source_url),
        status="ready_for_review",
        match_score=score,
        matched_keywords_json=json.dumps(matched_keywords, ensure_ascii=False),
        answers_json="[]",
        resume_filename=profile.resume_original_filename,
        notes="Prepared by AI Job Hunter from a configured public job source. Review all details before submitting.",
    )
    default_questions = [
        f"Tell us about your experience relevant to the {job.title} role.",
        f"Why are you interested in the {job.title} role at {job.company}?",
        "Which relevant skills make you a good fit for this position?",
        "What is your notice period?",
    ]
    app.answers_json = json.dumps(draft_screening_answers(profile, app, default_questions), ensure_ascii=False)
    return app


@app.post("/api/applications/queue", response_model=ApplicationQueueResponse, tags=["applications"])
def queue_selected_applications(payload: ApplicationQueueRequest, db: Session = Depends(get_db)):
    """Create or reuse review records for selected leads; never submit to employers."""
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        profile = CandidateProfile(id=1)
        db.add(profile)
        db.flush()

    items: list[ApplicationQueueItem] = []
    created_count = 0
    reused_count = 0
    saved_jobs_created = 0
    failed_count = 0
    seen_urls: set[str] = set()

    for selected in payload.jobs:
        title = selected.title.strip()
        company = selected.company.strip()
        try:
            canonical_url = canonicalize_job_url(str(selected.source_url))
            if canonical_url in seen_urls:
                reused_count += 1
                existing = db.scalar(select(Application).where(Application.source_url == canonical_url))
                items.append(ApplicationQueueItem(
                    job_id=selected.job_id, application_id=existing.id if existing else None,
                    title=title, company=company, action="reused",
                    message="Duplicate URL within this batch; only one application record is kept.",
                ))
                continue
            seen_urls.add(canonical_url)

            job = db.get(Job, selected.job_id) if selected.job_id is not None else None
            if job is not None and canonicalize_job_url(job.source_url) != canonical_url:
                raise ValueError("The selected job ID does not match the supplied source URL.")
            if job is None:
                job = db.scalar(select(Job).where(Job.source_url == canonical_url))
            if job is None:
                # Unverified-date leads are allowed into the review queue, but remain unverified.
                job = Job(
                    title=title[:250], company=company[:250],
                    location=(selected.location or "Not specified")[:250],
                    work_mode=(selected.work_mode or "Unknown")[:50],
                    source_url=canonical_url,
                    description=(selected.description or "")[:20000],
                    posted_at=selected.posted_at,
                    verified_posting_time=bool(selected.verified_posting_time),
                )
                db.add(job)
                db.flush()
                saved_jobs_created += 1

            existing_app = db.scalar(select(Application).where(Application.source_url == canonical_url))
            if existing_app is not None:
                reused_count += 1
                items.append(ApplicationQueueItem(
                    job_id=job.id, application_id=existing_app.id, title=existing_app.title,
                    company=existing_app.company, action="reused",
                    message=f"Already tracked as '{existing_app.status.replace('_', ' ')}'; existing answers and status were preserved.",
                ))
                db.commit()
                continue

            app_record = _application_for_job(
                job, int(selected.match_score), [str(k)[:100] for k in selected.matched_keywords[:30]], profile
            )
            app_record.notes = (
                "Added to the application review queue by the user. Screening answers are drafts; "
                "no application has been submitted. Verify the original listing and review every answer."
            )
            db.add(app_record)
            try:
                db.commit()
                db.refresh(app_record)
            except IntegrityError:
                db.rollback()
                existing_app = db.scalar(select(Application).where(Application.source_url == canonical_url))
                if existing_app is not None:
                    reused_count += 1
                    items.append(ApplicationQueueItem(
                        job_id=job.id, application_id=existing_app.id, title=existing_app.title,
                        company=existing_app.company, action="reused",
                        message="Another request already added this job; the existing record was reused.",
                    ))
                else:
                    failed_count += 1
                    items.append(ApplicationQueueItem(
                        job_id=selected.job_id, application_id=None, title=title, company=company,
                        action="failed", message="Could not save this application record. Retry this job individually.",
                    ))
                continue

            created_count += 1
            items.append(ApplicationQueueItem(
                job_id=job.id, application_id=app_record.id, title=app_record.title,
                company=app_record.company, action="created",
                message="Added to Application History with profile-grounded answer drafts. Not submitted.",
            ))
        except Exception as exc:
            db.rollback()
            failed_count += 1
            logger.warning("Application queue failed for %s at %s: %s", title, company, exc)
            items.append(ApplicationQueueItem(
                job_id=selected.job_id, application_id=None, title=title, company=company,
                action="failed", message=f"Could not queue this job ({type(exc).__name__}). Check the listing URL and retry.",
            ))

    requested = len(payload.jobs)
    return ApplicationQueueResponse(
        requested=requested, created=created_count, reused=reused_count,
        saved_jobs_created=saved_jobs_created, failed=failed_count, submitted=0, items=items,
        message=(f"Application queue updated: {created_count} created, {reused_count} already tracked, "
                 f"{failed_count} failed. No applications were submitted."),
    )


@app.get("/api/applications", response_model=list[ApplicationRead], tags=["applications"])
def list_applications(
    limit: int = Query(default=200, ge=1, le=1000),
    status: str | None = Query(default=None, max_length=40),
    db: Session = Depends(get_db),
):
    statement = select(Application).order_by(Application.created_at.desc()).limit(limit)
    if status:
        statement = select(Application).where(Application.status == status.lower()).order_by(Application.created_at.desc()).limit(limit)
    return db.scalars(statement).all()


@app.post("/api/applications", response_model=ApplicationRead, status_code=201, tags=["applications"])
def create_application(payload: ApplicationCreate, db: Session = Depends(get_db)):
    source_url = canonicalize_job_url(str(payload.source_url))
    existing = db.scalar(select(Application).where(Application.source_url == source_url))
    if existing:
        raise HTTPException(status_code=409, detail="This vacancy already exists in your application history.")
    values = payload.model_dump()
    values["source_url"] = source_url
    values["platform"] = platform_from_url(source_url) if payload.platform == "unknown" else payload.platform
    app_record = Application(**values)
    db.add(app_record)
    try:
        db.commit()
        db.refresh(app_record)
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="This vacancy already exists in your application history.")
    return app_record


@app.post("/api/applications/from-job/{job_id}", response_model=ApplicationRead, status_code=201, tags=["applications"])
def track_application_from_job(job_id: int, db: Session = Depends(get_db)):
    job = db.get(Job, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Saved job not found.")
    canonical_url = canonicalize_job_url(job.source_url)
    existing = db.scalar(select(Application).where(Application.source_url == canonical_url))
    if existing:
        raise HTTPException(status_code=409, detail="This vacancy is already in application history.")
    match = next((item for item in _get_rule_job_matches(limit=200, db=db) if item.job.id == job.id), None)
    profile = db.get(CandidateProfile, 1) or CandidateProfile(id=1)
    app_record = _application_for_job(job, match.match_score if match else 0, match.matched_keywords if match else [], profile)
    db.add(app_record)
    db.commit()
    db.refresh(app_record)
    return app_record


@app.patch("/api/applications/{application_id}", response_model=ApplicationRead, tags=["applications"])
def update_application(application_id: int, payload: ApplicationUpdate, db: Session = Depends(get_db)):
    app_record = db.get(Application, application_id)
    if app_record is None:
        raise HTTPException(status_code=404, detail="Application record not found.")
    changes = payload.model_dump(exclude_unset=True)
    for key, value in changes.items():
        if value is not None:
            setattr(app_record, key, value)
    if app_record.status in {"submitted", "submitted_unconfirmed"} and app_record.applied_at is None:
        app_record.applied_at = _utc_now()
        app_record.last_attempt_at = app_record.applied_at
    db.commit()
    db.refresh(app_record)
    return app_record


@app.post("/api/applications/{application_id}/answers/generate", response_model=ApplicationRead, tags=["applications"])
async def generate_application_answers(application_id: int, payload: ScreeningAnswerRequest, db: Session = Depends(get_db)):
    app_record = db.get(Application, application_id)
    if app_record is None:
        raise HTTPException(status_code=404, detail="Application record not found.")
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        raise HTTPException(status_code=400, detail="Complete your candidate profile first.")
    answers = await draft_screening_answers_with_ai(profile, app_record, payload.questions)
    try:
        previous_answers = json.loads(app_record.answers_json or "[]")
    except (ValueError, TypeError):
        previous_answers = []
    previous_by_question = {
        " ".join(re.sub(r"[^a-z0-9]+", " ", str(item.get("question") or "").lower()).split()): item
        for item in previous_answers if isinstance(item, dict)
    } if isinstance(previous_answers, list) else {}
    merged = dict(previous_by_question)
    for answer in answers:
        key = " ".join(re.sub(r"[^a-z0-9]+", " ", str(answer.get("question") or "").lower()).split())
        previous = merged.get(key)
        if previous and previous.get("status") == "verified":
            continue
        if previous and str(previous.get("source") or "").startswith("User edited") and previous.get("answer"):
            continue
        merged[key] = answer
    app_record.answers_json = json.dumps(list(merged.values()), ensure_ascii=False)
    if app_record.status == "draft":
        app_record.status = "ready_for_review"
    app_record.last_attempt_at = _utc_now()
    app_record.submission_message = "Screening answers drafted from profile/resume data. Review answers marked needs_review before use."
    db.commit()
    db.refresh(app_record)
    return app_record


@app.patch("/api/applications/{application_id}/answers/{answer_index}", response_model=ApplicationRead, tags=["applications"])
def review_screening_answer(
    application_id: int,
    answer_index: int,
    payload: ScreeningAnswerUpdate,
    db: Session = Depends(get_db),
):
    app_record = db.get(Application, application_id)
    if app_record is None:
        raise HTTPException(status_code=404, detail="Application record not found.")
    try:
        answers = json.loads(app_record.answers_json or "[]")
    except (ValueError, TypeError):
        answers = []
    if not isinstance(answers, list) or answer_index < 0 or answer_index >= len(answers):
        raise HTTPException(status_code=404, detail="Screening answer not found.")
    answer = answers[answer_index]
    if not isinstance(answer, dict):
        raise HTTPException(status_code=422, detail="Screening answer record is invalid.")
    answer["answer"] = payload.answer
    answer["status"] = payload.status
    answer["source"] = "User reviewed and confirmed" if payload.status == "verified" else "User edited; review still required"
    answer["confidence"] = "high" if payload.status == "verified" else "medium"
    app_record.answers_json = json.dumps(answers, ensure_ascii=False)
    app_record.last_attempt_at = _utc_now()
    db.commit()
    db.refresh(app_record)
    return app_record


@app.get("/api/automation/settings", response_model=AutomationSettingsResponse, tags=["automation"])
def automation_settings():
    return AutomationSettingsResponse(
        auto_submit_enabled=settings.auto_submit_enabled,
        supported_hosts=SUPPORTED_APPLICATION_HOSTS,
        note=("Auto-submit is off by default. Browser automation is limited to supported Greenhouse and Lever-hosted application pages; "
              "the agent will not bypass CAPTCHA, identity checks, or unsupported-site restrictions."),
    )


@app.get("/api/automation/diagnostics", tags=["automation"])
async def automation_diagnostics():
    """Explain whether the active backend environment can launch Playwright Chromium."""
    return await browser_runtime_diagnostics()


@app.post("/api/applications/{application_id}/browser-fill", response_model=BrowserApplyResponse, tags=["automation"])
async def browser_fill_application(application_id: int, payload: BrowserApplyRequest, db: Session = Depends(get_db)):
    app_record = db.get(Application, application_id)
    if app_record is None:
        raise HTTPException(status_code=404, detail="Application record not found.")
    profile = db.get(CandidateProfile, 1)
    if profile is None:
        raise HTTPException(status_code=400, detail="Complete your candidate profile first.")
    if payload.confirm_submit and not settings.auto_submit_enabled:
        raise HTTPException(status_code=403, detail="Submission is disabled. Set AUTO_SUBMIT_ENABLED=true in backend/.env and restart the backend only after reviewing the risks and site rules.")
    try:
        result = await fill_application_form(app_record, profile, RESUME_DIR, confirm_submit=payload.confirm_submit)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except Exception as exc:
        app_record.status = "failed"
        app_record.last_attempt_at = _utc_now()
        app_record.submission_message = f"Browser automation error: {str(exc)[:500]}"
        db.commit()
        raise HTTPException(status_code=502, detail=f"Browser automation failed: {str(exc)[:300]}") from exc

    app_record.status = result["status"]
    app_record.last_attempt_at = _utc_now()
    app_record.submission_message = result["message"]
    if result.get("submitted") or result.get("status") == "submitted_unconfirmed":
        app_record.applied_at = _utc_now()
    db.commit()
    return BrowserApplyResponse(application_id=app_record.id, **result)


@app.post("/api/agent/run", response_model=AgentRunResponse, tags=["agent"])
async def run_job_agent(payload: AgentRunRequest, db: Session = Depends(get_db)):
    started_at = _utc_now()
    # A requested auto-submit must never block useful discovery. If submission is
    # disabled in configuration, continue searching and preparing applications,
    # but explicitly skip the browser-submit stage and return a warning.
    auto_apply_allowed = bool(payload.auto_apply and settings.auto_submit_enabled)

    profile = db.get(CandidateProfile, 1)
    if profile is None or not profile.resume_text.strip():
        # The agent may still search, but answers will be sparse until a resume/profile is available.
        pass
    query = payload.query.strip() or (profile.target_titles if profile else "")
    discovery = await discover_jobs(
        DiscoveryRequest(
            query=query, hours=payload.hours, include_unverified=payload.include_unverified,
            region=payload.region, custom_location=payload.custom_location,
            work_modes=payload.work_modes, employment_types=payload.employment_types,
        ), db
    )
    if discovery.sources_checked == 0:
        no_source_warnings = ["No enabled public job sources are configured. Add company boards or public aggregators under Find Jobs. This searches configured sources, not every website on the internet."]
        if payload.auto_apply and not settings.auto_submit_enabled:
            no_source_warnings.append("Auto-submit was requested but is disabled by configuration (AUTO_SUBMIT_ENABLED=false). No applications were submitted; job discovery was allowed to finish.")
        return AgentRunResponse(
            started_at=started_at, finished_at=_utc_now(), query=query, hours=payload.hours, sources_checked=0,
            discovered_fresh=0, matched_jobs=0, applications_created=0, already_tracked=0,
            warnings=no_source_warnings,
        )

    fresh_ids = {item.id for item in discovery.fresh_jobs if item.id is not None}
    match_list = await get_job_matches(limit=200, db=db)
    matching = [item for item in match_list if item.job.id in fresh_ids and item.match_score >= payload.min_match_score]
    matching.sort(key=lambda item: item.match_score, reverse=True)
    candidates = matching[:payload.max_applications]
    created_ids: list[int] = []
    already_tracked = 0
    warnings: list[str] = []
    if payload.auto_apply and not settings.auto_submit_enabled:
        warnings.append("Auto-submit was requested but is disabled by configuration (AUTO_SUBMIT_ENABLED=false). Discovery and application preparation continued; no applications were submitted. Review settings and screening answers before enabling submission.")
    if profile is None:
        warnings.append("Candidate profile is missing. Add your profile and resume before preparing applications.")

    if payload.prepare_applications:
        for match in candidates:
            job = match.job
            canonical_url = canonicalize_job_url(job.source_url)
            existing = db.scalar(select(Application).where(Application.source_url == canonical_url))
            if existing:
                already_tracked += 1
                continue
            profile = profile or CandidateProfile(id=1)
            app_record = _application_for_job(job, match.match_score, match.matched_keywords, profile)
            db.add(app_record)
            try:
                db.flush()
            except IntegrityError:
                db.rollback()
                already_tracked += 1
                continue
            created_ids.append(app_record.id)

    db.commit()
    browser_attempted = 0
    submitted = 0
    needs_review = 0
    if auto_apply_allowed:
        for application_id in created_ids:
            app_record = db.get(Application, application_id)
            profile = db.get(CandidateProfile, 1)
            if not app_record or not profile:
                continue
            if not host_is_supported(app_record.source_url):
                app_record.status = "needs_manual_action"
                app_record.submission_message = "Unsupported application host. This role was saved; open the original URL and apply manually."
                needs_review += 1
                continue
            browser_attempted += 1
            try:
                result = await fill_application_form(app_record, profile, RESUME_DIR, confirm_submit=True)
                app_record.status = result["status"]
                app_record.last_attempt_at = _utc_now()
                app_record.submission_message = result["message"]
                if result.get("submitted") or result.get("status") == "submitted_unconfirmed":
                    app_record.applied_at = _utc_now()
                if result.get("submitted"):
                    submitted += 1
                else:
                    needs_review += 1
                    # A single human-review item pauses batch automation so its visible browser is not replaced.
                    warnings.append(f"Paused auto-apply after {app_record.company} — {app_record.title} needs review or manual completion.")
                    db.commit()
                    break
            except Exception as exc:
                app_record.status = "failed"
                app_record.last_attempt_at = _utc_now()
                app_record.submission_message = f"Automation failed: {str(exc)[:400]}"
                warnings.append(f"Could not process {app_record.company} — {app_record.title}: {str(exc)[:180]}")
                needs_review += 1
            db.commit()

    if not profile or not profile.resume_original_filename:
        warnings.append("No resume file is attached to the profile. Screening answers may be incomplete and automated forms cannot attach a resume.")
    warnings.append("Coverage includes enabled configured sources plus optional broad-web research when BRAVE_SEARCH_API_KEY is configured. Search results still need individual verification; unsupported application sites require a permitted connector or manual application.")
    return AgentRunResponse(
        started_at=started_at, finished_at=_utc_now(), query=query, hours=payload.hours, sources_checked=discovery.sources_checked,
        discovered_fresh=discovery.fresh_count, matched_jobs=len(matching), applications_created=len(created_ids),
        already_tracked=already_tracked, browser_attempted=browser_attempted, submitted=submitted,
        needs_review=needs_review, warnings=warnings, application_ids=created_ids,
    )


def _json_strings(value: str | None) -> list:
    try:
        result = json.loads(value or "[]")
        return result if isinstance(result, list) else []
    except (TypeError, ValueError):
        return []


def _excel_safe(value):
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


@app.get("/api/analytics", response_model=AnalyticsResponse, tags=["analytics"])
def application_analytics(db: Session = Depends(get_db)):
    from collections import Counter
    now = _utc_now()
    applications = db.scalars(select(Application)).all()
    by_status = Counter(app.status for app in applications)
    by_platform = Counter(app.platform or "unknown" for app in applications)
    def recent(dt, days):
        if not dt:
            return False
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt >= now - timedelta(days=days)
    submitted_statuses = {"submitted", "submitted_unconfirmed"}
    awaiting_statuses = {"draft", "ready_for_review", "needs_manual_action"}
    return AnalyticsResponse(
        total_applications=len(applications),
        submitted=sum(1 for app in applications if app.status in submitted_statuses or app.applied_at is not None),
        awaiting_review=sum(1 for app in applications if app.status in awaiting_statuses),
        interviews=by_status.get("interview", 0), offers=by_status.get("offer", 0), rejected=by_status.get("rejected", 0),
        applied_last_7_days=sum(1 for app in applications if recent(app.applied_at, 7)),
        applied_last_30_days=sum(1 for app in applications if recent(app.applied_at, 30)),
        by_status=dict(by_status), by_platform=dict(by_platform),
    )


@app.get("/api/analytics/export.xlsx", tags=["analytics"])
def export_application_history(db: Session = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="Excel export dependency missing. Run `pip install openpyxl`.") from exc

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Analytics"
    apps = db.scalars(select(Application).order_by(Application.created_at.desc())).all()
    summary.append(["AI Job Hunter — Application Analytics", "Value"])
    stats = application_analytics(db)
    rows = [
        ("Total tracked applications", stats.total_applications), ("Submitted / confirmed attempts", stats.submitted),
        ("Awaiting review", stats.awaiting_review), ("Interviews", stats.interviews), ("Offers", stats.offers),
        ("Rejected", stats.rejected), ("Applied in last 7 days", stats.applied_last_7_days),
        ("Applied in last 30 days", stats.applied_last_30_days), ("", ""), ("Status", "Count"),
    ]
    rows.extend(sorted(stats.by_status.items()))
    rows.extend([("", ""), ("Platform", "Count")])
    rows.extend(sorted(stats.by_platform.items()))
    for row in rows:
        summary.append([_excel_safe(value) for value in row])

    sheet = workbook.create_sheet("Applications")
    headers = ["Application ID", "Job title", "Company", "Location", "Work mode", "Platform", "Status", "Match score", "Job / apply URL", "Applied at (UTC)", "Created at (UTC)", "Updated at (UTC)", "Resume filename", "Matched keywords", "Notes", "Last attempt result", "Confirmation ID"]
    sheet.append(headers)
    answers_sheet = workbook.create_sheet("Screening Answers")
    answers_sheet.append(["Application ID", "Job title", "Company", "Question", "Draft answer", "Review status", "Confidence", "Source / evidence"])
    for app_record in apps:
        sheet.append([_excel_safe(x) for x in [
            app_record.id, app_record.title, app_record.company, app_record.location, app_record.work_mode,
            app_record.platform, app_record.status, app_record.match_score, app_record.source_url,
            app_record.applied_at.isoformat() if app_record.applied_at else "",
            app_record.created_at.isoformat() if app_record.created_at else "",
            app_record.updated_at.isoformat() if app_record.updated_at else "", app_record.resume_filename or "",
            ", ".join(str(x) for x in _json_strings(app_record.matched_keywords_json)), app_record.notes,
            app_record.submission_message, app_record.confirmation_id or "",
        ]])
        for answer in _json_strings(app_record.answers_json):
            if isinstance(answer, dict):
                answers_sheet.append([_excel_safe(x) for x in [app_record.id, app_record.title, app_record.company, answer.get("question", ""), answer.get("answer", ""), answer.get("status", ""), answer.get("confidence", ""), answer.get("source", "")]])

    header_fill = PatternFill("solid", fgColor="292B52")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in worksheet.columns:
            values = [len(str(cell.value or "")) for cell in list(column_cells)[:100]]
            width = min(48, max(12, max(values or [12]) + 2))
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"ai_job_hunter_application_history_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
